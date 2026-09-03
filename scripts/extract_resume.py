# -*- coding: utf-8 -*-
"""
简历文本提取：把 word / pdf 简历转成纯文本（入库协议第一步）

用法：
    python extract_resume.py 简历.docx
    python extract_resume.py 简历.pdf

依赖：
    pip install python-docx pdfplumber
"""
import sys
import os


def extract_docx(path):
    try:
        from docx import Document
    except ImportError:
        sys.exit("缺少依赖：pip install python-docx")
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
    return "\n".join(lines)


def extract_pdf(path):
    try:
        import pdfplumber
    except ImportError:
        sys.exit("缺少依赖：pip install pdfplumber")
    text = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text.append(t)
    return "\n".join(text)


def main():
    if len(sys.argv) < 2:
        sys.exit("用法: python extract_resume.py 简历.docx 或 简历.pdf")
    path = sys.argv[1]
    if not os.path.exists(path):
        sys.exit(f"文件不存在: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        text = extract_docx(path)
    elif ext == ".pdf":
        text = extract_pdf(path)
    else:
        sys.exit("仅支持 .docx 和 .pdf")
    out = os.path.splitext(path)[0] + ".md"
    with open(out, "w", encoding="utf-8") as f:
        f.write(text)
    print(f"已提取到: {out}")


if __name__ == "__main__":
    main()
